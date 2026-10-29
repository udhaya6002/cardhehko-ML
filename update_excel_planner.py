import openpyxl

wb = openpyxl.load_workbook('used-car-ml-master-planner.xlsx')

# 1. Update Master Plan sheet
ws_plan = wb['Master Plan']
updated_tasks = 0
for r in range(7, 39):
    task_id = ws_plan.cell(r, 1).value
    if task_id:
        ws_plan.cell(r, 13, value="Done")
        updated_tasks += 1
print(f"Updated {updated_tasks} tasks to 'Done' in Master Plan.")

# 2. Update Deliverables sheet
ws_deliv = wb['Deliverables']
updated_deliv = 0
for r in range(5, 25):
    item = ws_deliv.cell(r, 1).value
    if item:
        ws_deliv.cell(r, 4, value="Done")
        updated_deliv += 1
print(f"Updated {updated_deliv} items to 'Done' in Deliverables.")

wb.save('used-car-ml-master-planner.xlsx')
print("Successfully saved updated used-car-ml-master-planner.xlsx!")
