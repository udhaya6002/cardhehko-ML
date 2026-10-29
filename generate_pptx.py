import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
NAVY = RGBColor(26, 35, 126)
TEAL = RGBColor(0, 150, 136)
DARK_GRAY = RGBColor(40, 40, 40)
LIGHT_GRAY = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
CORAL = RGBColor(231, 76, 60)
BLUE = RGBColor(33, 150, 243)

def add_header(slide, title_text, category="USED CAR PRICE PREDICTION | ML COURSE PROJECT"):
    # Header background strip
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()

    # Category tracker
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.3))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category.upper()
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = TEAL

    # Main Title
    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.5), Inches(0.6))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = WHITE

blank_layout = prs.slide_layouts[6]

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()

# Decorative accent bar
bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.15), Inches(2.8))
bar.fill.solid()
bar.fill.fore_color.rgb = TEAL
bar.line.fill.background()

tb = slide1.shapes.add_textbox(Inches(1.6), Inches(2.0), Inches(10.5), Inches(3.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Used Car Price Prediction"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = WHITE

p2 = tf.add_paragraph()
p2.text = "A Leakage-Free Machine Learning Regression Framework"
p2.font.size = Pt(22)
p2.font.color.rgb = TEAL
p2.space_before = Pt(10)

p3 = tf.add_paragraph()
p3.text = "Featuring Regex Parsing, Domain Feature Engineering, 6-Algorithm CV Benchmarking, and Segment Error Profiling"
p3.font.size = Pt(14)
p3.font.color.rgb = RGBColor(200, 220, 240)
p3.space_before = Pt(15)

p4 = tf.add_paragraph()
p4.text = "Vellore Institute of Technology, Andhra Pradesh (VIT-AP) | Machine Learning Course Project"
p4.font.size = Pt(12)
p4.font.color.rgb = RGBColor(160, 180, 200)
p4.space_before = Pt(30)

# ==============================================================================
# SLIDE 2: PROBLEM STATEMENT & MOTIVATION
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Problem Statement & Market Motivation", "Executive Overview")

tb = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "The Valuation Challenge"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = NAVY

bullets = [
    "Used car pricing suffers from significant information asymmetry between buyers and sellers.",
    "Manual appraisals are subjective, time-intensive, and prone to individual dealer biases.",
    "Vehicle depreciation exhibits strong non-linearities: rapid drops in early years followed by plateauing decay.",
    "Complex feature interdependencies: brand prestige, mileage decay, fuel efficiency, engine output, and odometer wear."
]
for b in bullets:
    p = tf.add_paragraph()
    p.text = "- " + b
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(12)

# Right Box: Motivation
tb_r = slide2.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.5), Inches(5.4))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

p_r = tf_r.paragraphs[0]
p_r.text = "Business & Technical Motivation"
p_r.font.size = Pt(18)
p_r.font.bold = True
p_r.font.color.rgb = NAVY

bullets_r = [
    "Algorithmic Transparency: Provide instant, objective market valuations for resale platforms (CarDekho, Spinny, Cars24).",
    "Asset Underwriting: Deliver reliable residual asset value estimates for auto-lenders and insurance underwriters.",
    "Methodological Rigor: Eliminate pervasive data leakage found in academic literature through strict split isolation.",
    "Empirical Explainability: Understand exact monetary impact of mileage, vehicle age, and mechanical specifications."
]
for b in bullets_r:
    p = tf_r.add_paragraph()
    p.text = "- " + b
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(12)

# ==============================================================================
# SLIDE 3: OBJECTIVES & METHODOLOGICAL DISCIPLINE
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Core Project Objectives & Anti-Leakage Rules", "Methodology")

tb = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Key Project Objectives"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = NAVY

objs = [
    ("Strict Data Hygiene", "Split dataset 80/20 before any outlier capping, imputation, or scaling to eliminate data leakage."),
    ("Domain Feature Extraction", "Parse numeric units (kmpl, CC, bhp, Nm, rpm) and create depreciation rate metrics (car_age, km_per_year)."),
    ("Multi-Algorithmic Benchmark", "Train and tune 6 regression models across Linear, Regularized, Decision Tree, Bagging, and Boosting families."),
    ("Dual-Scale Valuation Tracking", "Evaluate all models on both log-transformed scale (variance stability) and real-world INR rupee scale."),
    ("Ablation & Segmentation", "Quantify exact predictive lift of engineered features and isolate prediction error across price and brand tiers.")
]
for title, desc in objs:
    p = tf.add_paragraph()
    p.text = f"{title}: "
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE
    p.space_before = Pt(12)
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = DARK_GRAY

# ==============================================================================
# SLIDE 4: DATASET OVERVIEW & TARGET DISTRIBUTION
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Dataset Overview: CarDekho Used Car Corpus", "Data Exploration")

tb = slide4.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.0), Inches(5.5))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Corpus Characteristics"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = NAVY

details = [
    "Raw Dataset: 8,128 records with 13 original attributes.",
    "De-duplication: Identified and pruned 1,679 duplicate listings (6,449 unique records retained).",
    "Target Variable: selling_price (INR)",
    "  - Min: Rs. 29,999 | Max: Rs. 10,000,000",
    "  - Mean: Rs. 638,272 | Median: Rs. 450,000",
    "Target Transformation: Extreme right-skewness corrected via log1p(y) mapping to stabilize variance."
]
for d in details:
    p = tf.add_paragraph()
    p.text = d
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(8)

if os.path.exists("figures/eda_price_distribution.png"):
    slide4.shapes.add_picture("figures/eda_price_distribution.png", Inches(6.0), Inches(1.5), width=Inches(6.8))

# ==============================================================================
# SLIDE 5: EXPLORATORY DATA ANALYSIS (CATEGORICAL)
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "Price Segmentation by Category & Fuel Type", "Data Exploration")

tb = slide5.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.2))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Key EDA Insights: Diesel vehicles command a premium over Petrol due to fuel economy and commercial utility; Automatic transmissions exhibit significantly higher median values than manuals; First owners retain the highest valuation."
p.font.size = Pt(13)
p.font.color.rgb = DARK_GRAY

if os.path.exists("figures/eda_price_by_category.png"):
    slide5.shapes.add_picture("figures/eda_price_by_category.png", Inches(0.8), Inches(2.7), width=Inches(11.7))

# ==============================================================================
# SLIDE 6: CORRELATION & FEATURE RELATIONSHIPS
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "Feature Relationships & Correlation Analysis", "Data Exploration")

tb = slide6.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.0), Inches(5.5))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Correlation Takeaways"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = NAVY

corr_notes = [
    "Manufacturing Year vs Price: Strong positive correlation (+0.41); newer vehicles retain high value.",
    "Kilometers Driven: Modest negative correlation (-0.23); high odometer readings discount prices non-linearly.",
    "Engine & Power: High positive correlation (+0.50 to +0.70) with selling price; larger displacements command premiums.",
    "Multicollinearity: Engine CC and Max Power show strong collinearity, requiring regularization or tree-based splits."
]
for c in corr_notes:
    p = tf.add_paragraph()
    p.text = "- " + c
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(10)

if os.path.exists("figures/eda_correlation_heatmap.png"):
    slide6.shapes.add_picture("figures/eda_correlation_heatmap.png", Inches(6.0), Inches(1.5), width=Inches(6.8))

# ==============================================================================
# SLIDE 7: FEATURE ENGINEERING & EXTRACTION
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
add_header(slide7, "Regex String Parsing & Domain Feature Engineering", "Feature Pipeline")

tb = slide7.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Derived & Extracted Feature Set"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = NAVY

feats = [
    ("mileage -> mileage_kmpl", "Regex extraction of numeric efficiency from strings (e.g., '23.4 kmpl', '18.9 km/kg')."),
    ("engine -> engine_cc", "Displacement volume parsed from CC strings (e.g., '1248 CC')."),
    ("max_power -> max_power_bhp", "Brake horsepower extracted as clean float (e.g., '103.52 bhp')."),
    ("torque -> torque_nm & max_torque_rpm", "Automated parsing handling units, kgm-to-Nm conversion (x9.80665), and RPM range extractions."),
    ("name -> brand", "Manufacturer brand token extracted from vehicle name (Maruti, Hyundai, BMW, Toyota, etc.)."),
    ("year -> car_age", "Calculated depreciation span: car_age = 2024 - year."),
    ("km_driven & car_age -> km_per_year", "Annual usage density metric capturing odometer wear rate per year of ownership.")
]
for name, desc in feats:
    p = tf.add_paragraph()
    p.text = f"{name}: "
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLUE
    p.space_before = Pt(8)
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = DARK_GRAY

# ==============================================================================
# SLIDE 8: SYSTEM ARCHITECTURE
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
add_header(slide8, "System Architecture: End-to-End Leakage-Free Pipeline", "Architecture")

if os.path.exists("figures/architecture.png"):
    slide8.shapes.add_picture("figures/architecture.png", Inches(0.8), Inches(1.4), width=Inches(11.7))

# ==============================================================================
# SLIDE 9: MODEL BENCHMARK TABLE
# ==============================================================================
slide9 = prs.slides.add_slide(blank_layout)
add_header(slide9, "Benchmark Results Across 6 Regression Algorithms", "Model Evaluation")

tb = slide9.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.0))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Comprehensive evaluation using 5-Fold Cross Validation on the training split, evaluated on unseen test data across both log1p and Indian Rupee (INR) scales."
p.font.size = Pt(13)
p.font.color.rgb = DARK_GRAY

# Add Table
rows = 7
cols = 7
left = Inches(0.8)
top = Inches(2.6)
width = Inches(11.7)
height = Inches(4.0)

table_shape = slide9.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

headers = ["Model", "CV R2", "Test R2 (Log)", "Test RMSE (Log)", "Test R2 (INR)", "Test RMSE (INR)", "Test MAE (INR)"]
for col_idx, h in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

data = [
    ["Gradient Boosting", "0.9150", "0.9132", "0.2201", "0.8987", "Rs. 149,075", "Rs. 73,299"],
    ["Random Forest", "0.9135", "0.9067", "0.2283", "0.8866", "Rs. 157,688", "Rs. 73,933"],
    ["Linear Regression", "0.8894", "0.8835", "0.2551", "0.8852", "Rs. 158,685", "Rs. 89,022"],
    ["Ridge Regression", "0.8894", "0.8835", "0.2550", "0.8829", "Rs. 160,267", "Rs. 89,101"],
    ["Lasso Regression", "0.8895", "0.8833", "0.2554", "0.8756", "Rs. 165,186", "Rs. 89,576"],
    ["Decision Tree", "0.8806", "0.8765", "0.2627", "0.8567", "Rs. 177,300", "Rs. 88,422"]
]

for row_idx, row_data in enumerate(data):
    for col_idx, val in enumerate(row_data):
        cell = table.cell(row_idx + 1, col_idx)
        cell.text = val
        if row_idx == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(232, 245, 233)
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.color.rgb = DARK_GRAY
            if row_idx == 0:
                p.font.bold = True
            p.alignment = PP_ALIGN.CENTER

# ==============================================================================
# SLIDE 10: MODEL COMPARISON VISUALIZATION
# ==============================================================================
slide10 = prs.slides.add_slide(blank_layout)
add_header(slide10, "Model Comparison: R2, RMSE & MAE Profiles", "Model Evaluation")

if os.path.exists("figures/eval_model_comparison.png"):
    slide10.shapes.add_picture("figures/eval_model_comparison.png", Inches(0.8), Inches(1.5), width=Inches(11.7))

# ==============================================================================
# SLIDE 11: CHAMPION MODEL DIAGNOSTICS (GRADIENT BOOSTING)
# ==============================================================================
slide11 = prs.slides.add_slide(blank_layout)
add_header(slide11, "Champion Model Diagnostics: Gradient Boosting", "Diagnostics")

if os.path.exists("figures/eval_gradient_boosting.png"):
    slide11.shapes.add_picture("figures/eval_gradient_boosting.png", Inches(0.8), Inches(1.5), width=Inches(11.7))

# ==============================================================================
# SLIDE 12: FEATURE IMPORTANCE & INTERPRETABILITY
# ==============================================================================
slide12 = prs.slides.add_slide(blank_layout)
add_header(slide12, "Model Interpretability: Key Price Drivers", "Interpretability")

tb = slide12.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.0), Inches(5.5))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Key Driver Insights"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = NAVY

insights = [
    "Maximum Power (bhp): Single highest predictor of vehicle price, capturing vehicle performance class.",
    "Car Age: Leading negative price determinant; confirms exponential initial depreciation.",
    "Engine CC: Strong contributor reflecting vehicle size and prestige tier.",
    "Brand Prestige: German luxury brands (BMW, Mercedes-Benz, Audi) add substantial price premiums.",
    "Transmission: Automatic transmissions command durable resale margins."
]
for ins in insights:
    p = tf.add_paragraph()
    p.text = "- " + ins
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(10)

if os.path.exists("figures/interp_gradient_boosting_importance.png"):
    slide12.shapes.add_picture("figures/interp_gradient_boosting_importance.png", Inches(6.0), Inches(1.5), width=Inches(6.8))

# ==============================================================================
# SLIDE 13: FEATURE ABLATION STUDY
# ==============================================================================
slide13 = prs.slides.add_slide(blank_layout)
add_header(slide13, "Feature Ablation: Quantifying Engineering Value", "Validation")

tb = slide13.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Ablation Experiment: Removing car_age and km_per_year"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = NAVY

bullets_abl = [
    "Hypothesis: Raw year and mileage alone are insufficient; explicit age and annual wear metrics improve predictive resolution.",
    "Full Pipeline Model Test R2: 0.9132 | RMSE: 0.2201",
    "Ablated Model (No car_age & km_per_year) Test R2: 0.8704 | RMSE: 0.2690",
    "Performance Impact: Delta R2 = -0.0428 | RMSE increases by +22.2%",
    "Conclusion: Derived depreciation features provide statistically significant explanatory power and are indispensable."
]
for b in bullets_abl:
    p = tf.add_paragraph()
    p.text = "- " + b
    p.font.size = Pt(14)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(14)

# ==============================================================================
# SLIDE 14: ERROR ANALYSIS & SEGMENT PROFILING
# ==============================================================================
slide14 = prs.slides.add_slide(blank_layout)
add_header(slide14, "Error Profiling Across Vehicle Price & Age Tiers", "Diagnostics")

tb = slide14.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Segment Error Findings (results/error_analysis.csv)"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = NAVY

err_points = [
    ("Budget Segment (< Rs. 205,000)", "Median absolute error is just Rs. 24,651 with Mean Absolute Percentage Error (MAPE) of 25.9%."),
    ("Midrange Segment (Rs. 350,000 - Rs. 700,000)", "Lowest relative error with MAPE of 13.0% - 13.5% and median error of Rs. 48,000 - Rs. 61,000."),
    ("Luxury Segment (> Rs. 700,000)", "Higher absolute error (median Rs. 103,100) but consistent percentage error (13.3%), validating model stability across tiers."),
    ("Brand Accuracy", "Top volume brands Maruti (N=419, median error Rs. 30,706) and Hyundai (N=271, median error Rs. 42,484) exhibit tight error bounds."),
    ("Vehicle Age Tiers", "Vehicles aged 3-6 years exhibit the lowest MAPE (10.75%), while 10+ year vehicles show higher variance (21.4% MAPE) due to mechanical condition heterogeneity.")
]
for title, desc in err_points:
    p = tf.add_paragraph()
    p.text = f"{title}: "
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE
    p.space_before = Pt(10)
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = DARK_GRAY

# ==============================================================================
# SLIDE 15: CONCLUSION & DELIVERABLES SUMMARY
# ==============================================================================
slide15 = prs.slides.add_slide(blank_layout)
add_header(slide15, "Project Summary & Verified Deliverables", "Conclusion")

tb = slide15.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Summary of Contributions"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = NAVY

summaries = [
    "End-to-End Leakage-Free Pipeline: Built with scikit-learn, joblib, and pandas, executing all scaling and imputing strictly on training partitions.",
    "Superior Champion Model: Tuned Gradient Boosting achieved Test R2 = 0.9132 (Log) and 0.8987 (Rupee), with MAE of Rs. 73,299.",
    "Ablation Confirmed: Feature engineering provides +0.0428 boost in R2.",
    "Deliverables Complete: Runnable main.ipynb notebook, full 15-section report, literature review table, architecture diagrams, sample data, and 15-slide presentation deck."
]
for s in summaries:
    p = tf.add_paragraph()
    p.text = "- " + s
    p.font.size = Pt(14)
    p.font.color.rgb = DARK_GRAY
    p.space_before = Pt(12)

# Save presentation
prs.save("deliverables/presentation.pptx")
print("Saved deliverables/presentation.pptx successfully!")
